"""Local document extraction. One request per process; JSON lines carry progress/results."""
import sys, json, io, base64
from pathlib import Path
sys.path.insert(0, str(Path(__file__).resolve().parent))

def emit(value):
    print(json.dumps(value, ensure_ascii=False), flush=True)

def layout_pages(req):
    if Path(req['source']).suffix.lower() == '.docx':
        from document_docx import read_docx
        return read_docx(req['source'])
    from document_layout import recover, finish_pages, blocks_text, LAYOUT_VERSION
    from PIL import Image
    source = Path(req['source'])
    pages = req['pages']
    pdf = None
    if source.suffix.lower() == '.pdf':
        import pypdfium2
        pdf = pypdfium2.PdfDocument(str(source))
    try:
        for index, saved in enumerate(pages):
            if saved.get('method') != 'ocr' or saved.get('layoutVersion') == LAYOUT_VERSION or not saved.get('lines'):
                continue
            page = bitmap = None
            if pdf is not None:
                page = pdf[saved['page']-1]
                bitmap = page.render(scale=saved.get('renderScale', 2.5))
                picture = bitmap.to_pil().convert('RGB')
            else:
                picture = Image.open(source).convert('RGB')
            try:
                saved['blocks'] = recover(picture, saved['lines'])
                saved['layoutVersion'] = LAYOUT_VERSION
            finally:
                picture.close()
                if bitmap is not None: bitmap.close()
                if page is not None: page.close()
            emit({'type': 'progress', 'completed': index+1, 'total': len(pages)})
        finish_pages(pages)
        for saved in pages:
            if saved.get('blocks'):
                saved['structuredText'] = blocks_text(saved['blocks'])
        return {'pages': pages}
    finally:
        if pdf is not None: pdf.close()

def preview_page(req):
    from PIL import Image
    source = Path(req['source'])
    pdf = page = bitmap = None
    if source.suffix.lower() == '.pdf':
        import pypdfium2
        pdf = pypdfium2.PdfDocument(str(source))
        page = pdf[int(req['page'])-1]
        bitmap = page.render(scale=1.6)
        picture = bitmap.to_pil().convert('RGB')
    else:
        picture = Image.open(source).convert('RGB')
    try:
        picture.thumbnail((1800, 2400))
        output = io.BytesIO()
        picture.save(output, format='WEBP', quality=85)
        return {'base64': base64.b64encode(output.getvalue()).decode('ascii')}
    finally:
        picture.close()
        if bitmap is not None: bitmap.close()
        if page is not None: page.close()
        if pdf is not None: pdf.close()

def parse(req):
    source = Path(req['source'])
    suffix = source.suffix.lower()
    if suffix == '.docx':
        from document_docx import read_docx
        return read_docx(source)
    if suffix in ('.txt', '.md'):
        return {'pages': [{'page': 1, 'text': source.read_text(encoding='utf-8-sig'), 'method': 'text'}], 'needsOcr': False}
    pages = []
    missing = []
    if suffix == '.pdf':
        from pypdf import PdfReader
        reader = PdfReader(source)
        if reader.is_encrypted and not reader.decrypt(''):
            raise ValueError('PDF 需要密码，请上传已解密的副本')
        for index, page in enumerate(reader.pages):
            content = page.extract_text() or ''
            if not content.strip() or req.get('forceOcr'):
                missing.append(index)
            pages.append({'page': index + 1, 'text': content, 'method': 'text' if content.strip() else 'image'})
    else:
        pages = [{'page': 1, 'text': '', 'method': 'image'}]
        missing = [0]
    if missing and not req.get('ocr'):
        return {'pages': pages, 'needsOcr': True, 'scanPages': [n + 1 for n in missing]}
    if missing:
        from rapidocr import RapidOCR
        from PIL import Image
        runtime = Path(sys.executable).parent
        model_config = json.loads((runtime / 'models' / 'paths.json').read_text(encoding='utf-8'))
        params = {key: str(runtime / 'models' / value) for key, value in model_config.items()}
        engine = RapidOCR(params=params)
        pdf = None
        if suffix == '.pdf':
            import pypdfium2
            pdf = pypdfium2.PdfDocument(str(source))
        try:
            for count, index in enumerate(missing):
                if pdf:
                    page = pdf[index]
                    bitmap = page.render(scale=2.5)
                    picture = bitmap.to_pil().convert('RGB')
                else:
                    picture = Image.open(source).convert('RGB')
                result = engine(picture)
                texts = getattr(result, 'txts', None)
                boxes = getattr(result, 'boxes', None)
                lines = []
                if texts is not None:
                    for n, value in enumerate(texts):
                        lines.append({'text': str(value), 'box': boxes[n].tolist() if boxes is not None else None})
                from document_layout import recover, LAYOUT_VERSION
                pages[index] = {'page': index + 1, 'text': '\n'.join(x['text'] for x in lines), 'method': 'ocr', 'lines': lines,
                                'blocks': recover(picture, lines), 'layoutVersion': LAYOUT_VERSION, 'renderScale': 2.5}
                picture.close()
                if pdf:
                    bitmap.close()
                    page.close()
                emit({'type': 'progress', 'completed': count + 1, 'total': len(missing), 'page': index + 1, 'pages': [pages[index]]})
        finally:
            if pdf:
                pdf.close()
    from document_layout import finish_pages, blocks_text
    finish_pages(pages)
    for page in pages:
        if page.get('blocks'): page['structuredText'] = blocks_text(page['blocks'])
    return {'pages': pages, 'needsOcr': False, 'scanPages': [n + 1 for n in missing]}

def template_info(req):
    import openpyxl
    wb = openpyxl.load_workbook(req['source'], read_only=False, data_only=False)
    result = []
    for ws in wb:
        indices = sorted(set(range(1, min(ws.max_row, 3) + 1)) | set(range(max(1, ws.max_row - 2), ws.max_row + 1)))
        result.append({'name': ws.title, 'rows': ws.max_row, 'columns': ws.max_column,
            'sampleRows': [{'row': row, 'values': [ws.cell(row, c).value for c in range(1, min(ws.max_column, 30) + 1)]} for row in indices],
            'merges': [str(r) for r in ws.merged_cells.ranges]})
    wb.close()
    return result

if __name__ == '__main__':
    try:
        request = json.loads(Path(sys.argv[1]).read_text(encoding='utf-8-sig'))
        operation = request.get('operation')
        value = template_info(request) if operation == 'template-info' else layout_pages(request) if operation == 'layout' else preview_page(request) if operation == 'preview' else parse(request)
        emit({'type': 'result', 'value': value})
    except Exception as exc:
        emit({'type': 'error', 'message': str(exc)})
        sys.exit(1)
